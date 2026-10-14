"""Сборка istok-audit/report.docx из catalog.md, findings.md, console.md, progress.md.

Запуск из корня репозитория:
    python3 istok-audit/tools/build_report.py
Промежуточный отчёт по пакету:
    python3 istok-audit/tools/build_report.py --sims 1-7 --out istok-audit/report-batch1.docx
"""

import argparse
import re
from collections import Counter
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

AUDIT = Path(__file__).resolve().parent.parent

CATEGORIES = ["Текст", "UX", "Теория", "Физика", "Ввод-вывод", "Техническая"]
SEVERITIES = ["Критическая", "Существенная", "Незначительная"]
FIELDS = [
    "Симуляция", "Категория", "Серьёзность", "Где", "Шаги", "Ожидалось",
    "Получено", "Расчёт", "Скриншот", "Уверенность",
]
FIELD_RE = re.compile(r"^- (%s):\s?(.*)$" % "|".join(FIELDS))
ID_RE = re.compile(r"^### (SIM-(\d+)-(\d+))\s*$")
COMMENT_RE = re.compile(r"<!--.*?-->", re.S)


def read(name):
    path = AUDIT / name
    return COMMENT_RE.sub("", path.read_text(encoding="utf-8")) if path.exists() else ""


def section(text, title):
    """Текст под заголовком '## title' до следующего заголовка '## '."""
    m = re.search(r"^## %s\s*$(.*?)(?=^## |\Z)" % re.escape(title), text, re.M | re.S)
    return m.group(1).strip() if m else ""


def md_table(text):
    rows = []
    for line in text.splitlines():
        line = line.strip()
        if not line.startswith("|") or re.match(r"^\|[\s:|-]+\|$", line):
            continue
        rows.append([c.strip() for c in line.strip("|").split("|")])
    return rows[1:] if rows else []


def parse_catalog():
    sims = {}
    for row in md_table(read("catalog.md")):
        if len(row) < 6 or not row[0].isdigit():
            continue
        num = int(row[0])
        sims[num] = dict(num=num, name=row[1], url=row[2], section=row[3],
                         audience=row[4], desc=row[5])
    return sims


def parse_findings():
    text = read("findings.md")
    body = re.split(r"^## Системные проблемы\s*$", text, flags=re.M)[0]
    findings, cur, key = [], None, None
    for line in body.splitlines():
        m = ID_RE.match(line)
        if m:
            cur = dict(id=m.group(1), sim=int(m.group(2)), n=int(m.group(3)))
            findings.append(cur)
            key = None
            continue
        if cur is None:
            continue
        if line.startswith("#"):
            cur, key = None, None
            continue
        fm = FIELD_RE.match(line)
        if fm:
            key = fm.group(1)
            cur[key] = fm.group(2).strip()
        elif key and line.strip():
            cur[key] = (cur[key] + "\n" + line.rstrip()).strip()
    systemic = [l.lstrip("-* ").strip()
                for l in section(text, "Системные проблемы").splitlines() if l.strip()]
    return findings, systemic


def parse_console():
    blocks = re.split(r"^## (SIM-\d+.*)$", read("console.md"), flags=re.M)
    return [(blocks[i].strip(), blocks[i + 1].strip()) for i in range(1, len(blocks), 2)]


def parse_range(spec):
    nums = set()
    for part in spec.split(","):
        a, _, b = part.partition("-")
        nums.update(range(int(a), int(b or a) + 1))
    return nums


def severity_rank(f):
    sev = f.get("Серьёзность", "")
    s = SEVERITIES.index(sev) if sev in SEVERITIES else len(SEVERITIES)
    unsure = "методист" in f.get("Уверенность", "")
    return (s, unsure, f["sim"], f["n"])


def add_table(doc, rows, header=True, widths=None):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(val)
            if header and r == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            if widths:
                cell.width = widths[c]
    return table


def add_finding(doc, f, fig_no):
    doc.add_heading(f["id"], level=3)
    rows = [[k, f[k]] for k in FIELDS[1:] if k not in ("Скриншот",) and f.get(k)]
    if rows:
        add_table(doc, rows, header=False, widths=[Cm(3.5), Cm(13)])
    shot = f.get("Скриншот", "")
    path = AUDIT / shot if shot else AUDIT / "screenshots" / (f["id"] + ".png")
    if path.exists():
        doc.add_picture(str(path), width=Cm(15))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph("Рис. %d — %s" % (fig_no, f["id"]), style="Caption")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return fig_no + 1
    p = doc.add_paragraph()
    p.add_run("Скриншот отсутствует: %s" % path.relative_to(AUDIT)).italic = True
    return fig_no


def build(out, sims_filter=None):
    sims = parse_catalog()
    findings, systemic = parse_findings()
    console = parse_console()
    progress = read("progress.md")
    env_rows = md_table(section(progress, "Окружение"))
    not_checked = section(progress, "Не удалось проверить")

    if sims_filter:
        sims = {k: v for k, v in sims.items() if k in sims_filter}
        findings = [f for f in findings if f["sim"] in sims_filter or f["sim"] == 0]
        console = [c for c in console
                   if int(re.match(r"SIM-(\d+)", c[0]).group(1)) in sims_filter]

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    # 1. Титул
    doc.add_heading("Аудит симуляций ИСТОК, edu-dev", level=0)
    doc.add_paragraph("Дата отчёта: %s" % date.today().isoformat())
    if sims_filter:
        doc.add_paragraph("Промежуточный отчёт, симуляции: %s"
                          % ", ".join(str(n) for n in sorted(sims_filter)))
    doc.add_paragraph("Симуляций в отчёте: %d" % len(sims))
    doc.add_paragraph("Находок: %d" % len(findings))
    if env_rows:
        doc.add_paragraph("Окружение:")
        add_table(doc, [["Параметр", "Значение"]] + env_rows)
    doc.add_page_break()

    # 2. Сводка
    doc.add_heading("Сводка", level=1)
    cnt = Counter((f.get("Категория"), f.get("Серьёзность")) for f in findings)
    rows = [["Категория"] + SEVERITIES + ["Всего"]]
    for cat in CATEGORIES:
        vals = [cnt[(cat, s)] for s in SEVERITIES]
        rows.append([cat] + vals + [sum(vals)])
    totals = [sum(r[i] for r in rows[1:]) for i in range(1, len(SEVERITIES) + 2)]
    rows.append(["Всего"] + totals)
    add_table(doc, rows)
    unknown = [f["id"] for f in findings
               if f.get("Категория") not in CATEGORIES or f.get("Серьёзность") not in SEVERITIES]
    if unknown:
        doc.add_paragraph("Не учтены в таблице (категория или серьёзность вне списка): "
                          + ", ".join(unknown))

    doc.add_heading("Главные находки", level=2)
    top = sorted(findings, key=severity_rank)[:10]
    if top:
        for f in top:
            doc.add_paragraph("%s — %s, %s. %s" % (
                f["id"], f.get("Серьёзность", "?"), f.get("Категория", "?"),
                f.get("Получено", "").split("\n")[0]), style="List Number")
    else:
        doc.add_paragraph("Находок нет.")

    doc.add_heading("Системные проблемы", level=2)
    if systemic:
        for s in systemic:
            doc.add_paragraph(s, style="List Bullet")
    else:
        doc.add_paragraph("Не выявлено.")
    doc.add_page_break()

    # 3. По симуляциям
    doc.add_heading("Находки по симуляциям", level=1)
    fig_no = 1
    by_sim = {}
    for f in findings:
        by_sim.setdefault(f["sim"], []).append(f)
    for num in [0] + sorted(sims):
        items = sorted(by_sim.get(num, []), key=lambda f: f["n"])
        if num == 0:
            if not items:
                continue
            doc.add_heading("SIM-00. Общие страницы сайта", level=2)
        else:
            s = sims[num]
            doc.add_heading("SIM-%02d. %s" % (num, s["name"]), level=2)
            doc.add_paragraph(s["url"])
            doc.add_paragraph(s["desc"])
        if not items:
            doc.add_paragraph("Находок нет.")
        for f in items:
            fig_no = add_finding(doc, f, fig_no)
    orphans = sorted(set(by_sim) - set(sims) - {0})
    if orphans:
        doc.add_paragraph("Находки с номером симуляции вне каталога: "
                          + ", ".join("SIM-%02d" % n for n in orphans))

    # 4. Приложения
    doc.add_page_break()
    doc.add_heading("Приложение А. Ошибки консоли", level=1)
    if console:
        for title, body in console:
            doc.add_heading(title, level=2)
            for line in body.splitlines():
                if line.strip():
                    doc.add_paragraph(line.lstrip("-* ").strip(), style="List Bullet")
    else:
        doc.add_paragraph("Записей нет.")

    doc.add_heading("Приложение Б. Что не удалось проверить", level=1)
    lines = [l.lstrip("-* ").strip() for l in not_checked.splitlines() if l.strip()]
    for line in lines or ["Ограничений не зафиксировано."]:
        doc.add_paragraph(line, style="List Bullet")

    doc.save(out)
    return len(sims), len(findings), fig_no - 1


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(AUDIT / "report.docx"))
    ap.add_argument("--sims", help="номера симуляций для промежуточного отчёта, напр. 1-7,9")
    args = ap.parse_args()
    n_sims, n_find, n_fig = build(args.out, parse_range(args.sims) if args.sims else None)
    print("%s: симуляций %d, находок %d, рисунков %d" % (args.out, n_sims, n_find, n_fig))


if __name__ == "__main__":
    main()
